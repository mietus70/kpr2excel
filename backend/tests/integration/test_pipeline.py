"""Integration: PDF -> SQLite -> corrections -> configured XLSX."""

from __future__ import annotations

import datetime as _dt
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fixtures.synthetic_kpir import DEFAULT_ROWS, build_synthetic_kpir
from kpir_converter.application.services import (
    ExportService,
    ExtractionService,
    ImportError_,
    ImportService,
)
from kpir_converter.domain.models import (
    ExportDefinition,
    ExportPolicy,
    ExportScope,
    MoneyRangeFilter,
    RowSelection,
    RowSelectionMode,
    ValueType,
)
from kpir_converter.infrastructure.pdf.adapter import PdfError


@pytest.fixture()
def imported(ctx, sample_pdf: Path):
    service = ImportService(ctx)
    batch_id = service.create_batch("Testowa paczka")
    document = service.import_stream(batch_id, "kpir-2018.pdf", iter([sample_pdf.read_bytes()]))
    ExtractionService(ctx).run(document.id)
    return batch_id, document.id


class TestImportAndExtract:
    def test_original_is_stored_under_uuid(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        row = ctx.documents.get(document_id)
        stored = Path(row["stored_path"])
        assert stored.exists()
        assert stored.name == "source.pdf"
        assert document_id in str(stored)
        # The uploaded name survives only as sanitised metadata.
        assert row["original_name"] == "kpir-2018.pdf"

    def test_records_are_persisted(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        assert ctx.records.count([document_id]) == len(DEFAULT_ROWS)

    def test_document_becomes_ready(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        assert ctx.documents.get(document_id)["status"] == "ready"

    def test_profile_and_period_detected(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        row = ctx.documents.get(document_id)
        assert row["detected_profile_id"] == "kpir_pl_2018@1"
        assert row["declared_period_from"] == "2018-01-01"

    def test_values_round_trip_through_sqlite(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        page = ctx.records.page_records([document_id], limit=10)
        first = page.records[0]
        assert first.cell("business_date").parsed_value_text == "2018-01-04"
        assert first.cell("total_income").parsed_value_text == "1300.50"
        # Money survives as exact decimal text, never as a float.
        assert Decimal(first.cell("total_income").parsed_value_text) == Decimal("1300.50")

    def test_source_spans_survive(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("business_date")
        assert cell.source is not None
        assert cell.source.page_number == 1
        assert 0 <= cell.source.bbox.x0 < cell.source.bbox.x1 <= 1

    def test_pagination_is_stable(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        first = ctx.records.page_records([document_id], limit=2)
        assert len(first.records) == 2
        assert first.next_cursor is not None
        second = ctx.records.page_records([document_id], cursor=first.next_cursor, limit=2)
        assert len(second.records) == 2
        assert {r.id for r in first.records}.isdisjoint({r.id for r in second.records})

    def test_streaming_iterator_returns_all(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        streamed = list(ctx.records.iter_records([document_id], chunk_size=2))
        assert len(streamed) == len(DEFAULT_ROWS)
        assert [r.logical_index for r in streamed] == sorted(r.logical_index for r in streamed)


class TestCorrections:
    def test_correction_overrides_extraction(self, ctx, imported, profile) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("contractor_name")

        applied, revision = ctx.records.apply_correction(
            cell.id,
            value_text="Poprawiona Nazwa",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,
            reason="literówka",
        )
        assert applied
        assert revision == cell.revision + 1

        updated = ctx.records.get_record(record.id).cell("contractor_name")
        assert updated.parsed_value_text == "Poprawiona Nazwa"
        assert updated.is_manual is True
        # The extraction value is still reachable for the UI.
        assert updated.extraction_meta["extracted_value"] == "Alfa Testowa Sp. z o.o."

    def test_stale_revision_is_rejected(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("notes")
        ctx.records.apply_correction(
            cell.id,
            value_text="A",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,
            reason=None,
        )
        applied, current = ctx.records.apply_correction(
            cell.id,
            value_text="B",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,  # stale
            reason=None,
        )
        assert applied is False
        assert current == cell.revision + 1

    def test_revert_restores_extraction(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("contractor_name")
        original = cell.parsed_value_text
        _ok, revision = ctx.records.apply_correction(
            cell.id,
            value_text="Inna",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,
            reason=None,
        )
        ok, _ = ctx.records.revert_correction(cell.id, revision)
        assert ok
        restored = ctx.records.get_record(record.id).cell("contractor_name")
        assert restored.parsed_value_text == original
        assert restored.is_manual is False

    def test_history_is_auditable(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("notes")
        _ok, rev = ctx.records.apply_correction(
            cell.id,
            value_text="X",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,
            reason="test",
        )
        ctx.records.revert_correction(cell.id, rev)
        history = ctx.records.cell_history(cell.id)
        assert [h["action"] for h in history] == ["revert", "set"]
        assert history[1]["reason"] == "test"

    def test_original_pdf_is_untouched(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        row = ctx.documents.get(document_id)
        import hashlib

        digest = hashlib.sha256(Path(row["stored_path"]).read_bytes()).hexdigest()
        assert digest == row["sha256"]


class TestConfiguredExport:
    def _definition(self, document_id: str, **kwargs) -> ExportDefinition:
        base = {
            "scope": ExportScope(type="documents", document_ids=(document_id,)),
            "column_keys": ("business_date", "evidence_number", "total_expenses"),
            "policy": ExportPolicy.DRAFT,
        }
        base.update(kwargs)
        return ExportDefinition(**base)

    def test_preview_counts_match_filter(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        service = ExportService(ctx)
        definition = self._definition(
            document_id,
            filters=(MoneyRangeFilter("total_expenses", min_text="400.00", max_text="1000.00"),),
        )
        preview = service.preview(definition)
        # Amounts are: none, 446.50, 3500.00, 89.10, 1000.00.
        # 446.50 and 1000.00 fall inside the inclusive 400..1000 range.
        assert preview.matching_row_count == 2
        # 3500.00 and 89.10 are outside the range.
        assert preview.excluded_by_range_count == 2
        # The row with no amount never satisfies an active money filter.
        assert preview.excluded_for_null_or_invalid_count == 1
        assert preview.column_count == 3
        assert preview.total_row_count == 5

    def test_export_produces_typed_xlsx(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        service = ExportService(ctx)
        definition = self._definition(document_id)
        export_id, _job = service.create(definition)
        path = service.build(export_id)

        workbook = load_workbook(path)
        sheet = workbook["KPiR"]
        assert sheet["A1"].value == "Data zdarzenia gospodarczego"
        assert sheet["C1"].value == "Razem wydatki"
        # Real Excel date, not a string.
        assert sheet["A2"].value == _dt.datetime(2018, 1, 4)
        # Numbers are numeric with the 0.00 format.
        expenses = [sheet.cell(row=r, column=3).value for r in range(2, 7)]
        assert Decimal(str(expenses[1])) == Decimal("446.50")
        assert sheet.cell(row=3, column=3).number_format == "0.00"
        # Empty stays empty, never zero.
        assert expenses[0] is None
        assert sheet.freeze_panes == "A2"

    def test_column_order_follows_user_choice(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        service = ExportService(ctx)
        definition = self._definition(
            document_id, column_keys=("total_expenses", "row_number", "business_date")
        )
        export_id, _ = service.create(definition)
        sheet = load_workbook(service.build(export_id))["KPiR"]
        assert [sheet.cell(row=1, column=c).value for c in (1, 2, 3)] == [
            "Razem wydatki",
            "Lp.",
            "Data zdarzenia gospodarczego",
        ]

    def test_filter_by_hidden_column(self, ctx, imported) -> None:
        """Filter on total_expenses but do not export it."""
        _batch_id, document_id = imported
        service = ExportService(ctx)
        definition = self._definition(
            document_id,
            column_keys=("evidence_number",),
            filters=(MoneyRangeFilter("total_expenses", min_text="1000.00"),),
        )
        export_id, _ = service.create(definition)
        sheet = load_workbook(service.build(export_id))["KPiR"]
        assert sheet.max_column == 1
        values = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
        # 3500.00 and 1000.00 pass the inclusive >= 1000.00 bound.
        assert values == ["LP/3/2018", "ZK/2018/5"]

    def test_explicit_row_selection(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        records = ctx.records.page_records([document_id], limit=10).records
        chosen = [records[0].id, records[4].id]
        service = ExportService(ctx)
        definition = self._definition(
            document_id,
            row_selection=RowSelection(
                mode=RowSelectionMode.EXPLICIT, included_record_ids=tuple(chosen)
            ),
        )
        export_id, _ = service.create(definition)
        sheet = load_workbook(service.build(export_id))["KPiR"]
        assert sheet.max_row == 3  # header + 2 rows

    def test_exclusions_in_all_matching(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        records = ctx.records.page_records([document_id], limit=10).records
        service = ExportService(ctx)
        definition = self._definition(
            document_id,
            row_selection=RowSelection(
                mode=RowSelectionMode.ALL_MATCHING, excluded_record_ids=(records[0].id,)
            ),
        )
        export_id, _ = service.create(definition)
        sheet = load_workbook(service.build(export_id))["KPiR"]
        assert sheet.max_row == 5  # header + 4 rows

    def test_export_uses_corrected_values(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=10).records[1]
        cell = record.cell("total_expenses")
        ctx.records.apply_correction(
            cell.id,
            value_text="9999.99",
            value_type=ValueType.MONEY,
            base_revision=cell.revision,
            reason="korekta",
        )
        service = ExportService(ctx)
        definition = self._definition(
            document_id,
            filters=(MoneyRangeFilter("total_expenses", min_text="9000.00"),),
        )
        export_id, _ = service.create(definition)
        sheet = load_workbook(service.build(export_id))["KPiR"]
        assert sheet.max_row == 2
        # openpyxl reads numeric cells back as float; the stored value keeps
        # full precision (verified as decimal text below).
        assert Decimal(str(sheet.cell(row=2, column=3).value)) == Decimal("9999.99")

    def test_snapshot_is_immutable(self, ctx, imported) -> None:
        """Changing data after creation must not change the frozen definition."""
        _batch_id, document_id = imported
        service = ExportService(ctx)
        export_id, _ = service.create(self._definition(document_id))
        stored = ctx.exports.get(export_id)["definition_json"]

        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("notes")
        ctx.records.apply_correction(
            cell.id,
            value_text="zmiana",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,
            reason=None,
        )
        assert ctx.exports.get(export_id)["definition_json"] == stored

    def test_stale_source_revision_conflicts(self, ctx, imported) -> None:
        from kpir_converter.domain.export import ExportValidationError

        _batch_id, document_id = imported
        service = ExportService(ctx)
        definition = self._definition(document_id, source_revision="obsolete-token")
        with pytest.raises(ExportValidationError) as exc:
            service.create(definition)
        assert exc.value.code == "STALE_PREVIEW"

    def test_formula_injection_is_neutralised(self, ctx, imported, profile) -> None:
        _batch_id, document_id = imported
        record = ctx.records.page_records([document_id], limit=1).records[0]
        cell = record.cell("event_description")
        ctx.records.apply_correction(
            cell.id,
            value_text="=SUM(A1:A9)",
            value_type=ValueType.TEXT,
            base_revision=cell.revision,
            reason=None,
        )
        service = ExportService(ctx)
        definition = self._definition(document_id, column_keys=("event_description",))
        export_id, _ = service.create(definition)
        sheet = load_workbook(service.build(export_id))["KPiR"]
        assert sheet.cell(row=2, column=1).value == "'=SUM(A1:A9)"

    def test_issues_sheet_optional(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        service = ExportService(ctx)
        export_id, _ = service.create(self._definition(document_id))
        assert load_workbook(service.build(export_id)).sheetnames == ["KPiR", "Metadane"]

        export_id2, _ = service.create(self._definition(document_id, include_issues_sheet=True))
        assert "Problemy" in load_workbook(service.build(export_id2)).sheetnames

    def test_export_records_hash_and_counts(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        service = ExportService(ctx)
        export_id, _ = service.create(self._definition(document_id))
        service.build(export_id)
        row = ctx.exports.get(export_id)
        assert row["status"] == "ready"
        assert row["result_row_count"] == 5
        assert row["result_column_count"] == 3
        assert len(row["sha256"]) == 64
        assert row["profile_id"] == "kpir_pl_2018"


class TestDeletion:
    def test_delete_removes_all_derived_data(self, ctx, imported) -> None:
        _batch_id, document_id = imported
        service = ExportService(ctx)
        export_id, _ = service.create(
            ExportDefinition(
                scope=ExportScope(type="documents", document_ids=(document_id,)),
                column_keys=("business_date",),
                policy=ExportPolicy.DRAFT,
            )
        )
        service.build(export_id)
        ExtractionService(ctx).render_page(document_id, 1)

        originals = ctx.settings.originals_dir / document_id
        pages = ctx.settings.pages_dir / document_id
        assert originals.exists() and pages.exists()

        ImportService(ctx).delete_document(document_id)

        assert not originals.exists()
        assert not pages.exists()
        assert ctx.documents.get(document_id) is None
        assert ctx.records.count([document_id]) == 0
        assert ctx.issues.counts_by_severity([document_id]) == {}
        assert ctx.exports.get(export_id) is None


class TestResilience:
    def test_encrypted_pdf_reports_error(self, ctx, tmp_path: Path) -> None:
        import pymupdf

        doc = pymupdf.open()
        doc.new_page()
        protected = tmp_path / "enc.pdf"
        doc.save(
            str(protected),
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw="owner",
            user_pw="user",
        )
        doc.close()

        service = ImportService(ctx)
        batch_id = service.create_batch("enc")
        document = service.import_stream(batch_id, "enc.pdf", iter([protected.read_bytes()]))
        with pytest.raises(PdfError) as exc:
            ExtractionService(ctx).run(document.id)
        assert exc.value.code == "PDF_ENCRYPTED"
        assert ctx.documents.get(document.id)["status"] == "failed"
        assert ctx.documents.get(document.id)["error_code"] == "PDF_ENCRYPTED"

    def test_pdf_without_text_requires_ocr(self, ctx, tmp_path: Path) -> None:
        import pymupdf

        doc = pymupdf.open()
        doc.new_page()  # blank page, no text layer
        blank = tmp_path / "scan.pdf"
        doc.save(str(blank))
        doc.close()

        service = ImportService(ctx)
        batch_id = service.create_batch("scan")
        document = service.import_stream(batch_id, "scan.pdf", iter([blank.read_bytes()]))
        with pytest.raises(ImportError_) as exc:
            ExtractionService(ctx).run(document.id)
        assert exc.value.code == "OCR_REQUIRED"
        rows, _ = ctx.issues.list([document.id], status="open")
        assert any(r["code"] == "OCR_REQUIRED" for r in rows)

    def test_non_pdf_is_rejected(self, ctx) -> None:
        service = ImportService(ctx)
        batch_id = service.create_batch("bad")
        with pytest.raises(ImportError_) as exc:
            service.import_stream(batch_id, "evil.pdf", iter([b"<html>nope</html>"]))
        assert exc.value.code == "NOT_A_PDF"

    def test_cancellation_stops_extraction(self, ctx, tmp_path: Path) -> None:
        pdf = build_synthetic_kpir(tmp_path / "many.pdf", rows_per_page=1)
        service = ImportService(ctx)
        batch_id = service.create_batch("cancel")
        document = service.import_stream(batch_id, "many.pdf", iter([pdf.read_bytes()]))

        calls = {"n": 0}

        def should_cancel() -> bool:
            calls["n"] += 1
            return calls["n"] > 2

        ExtractionService(ctx).run(document.id, should_cancel=should_cancel)
        # Pages processed before the cancel are kept as a partial, valid result.
        assert 0 < ctx.records.count([document.id]) < len(DEFAULT_ROWS)
