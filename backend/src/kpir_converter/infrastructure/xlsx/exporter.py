"""XLSX writer.

- dates are written as real Excel dates, amounts as numbers with format ``0.00``;
- an empty value stays an empty cell, never a zero;
- text cells are protected against formula injection (``=``, ``+``, ``-``, ``@``);
- the sheet is streamed (``write_only``) so a huge batch does not blow up RAM.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ...domain.models import KpirRecord
from ...domain.profiles import ColumnType, KpirProfile

__all__ = ["XlsxExportResult", "write_export", "sanitize_text"]

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass(slots=True)
class XlsxExportResult:
    path: Path
    row_count: int
    column_count: int
    sha256: str
    size_bytes: int


def sanitize_text(value: str) -> str:
    """Neutralise spreadsheet formula injection while keeping the text readable."""
    if value and value[0] in _FORMULA_PREFIXES:
        return "'" + value
    return value


def _as_decimal(text: str | None) -> Decimal | None:
    if text is None or text == "":
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _as_date(text: str | None) -> _dt.date | None:
    if not text:
        return None
    try:
        return _dt.date.fromisoformat(text)
    except ValueError:
        return None


def write_export(
    target: Path,
    profile: KpirProfile,
    column_keys: Sequence[str],
    records: Iterable[KpirRecord],
    *,
    issues: Sequence[dict] | None = None,
    include_issues_sheet: bool = False,
    metadata: dict[str, str] | None = None,
) -> XlsxExportResult:
    """Write the projected records to a workbook. Purely a projection of data."""
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=profile.sheet_name[:31])

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDE6F0")
    header_alignment = Alignment(vertical="center", wrap_text=True)

    columns = [profile.column(key) for key in column_keys]

    # In write-only mode sheet-level settings must be applied *before* the rows
    # are streamed out, otherwise they are silently dropped from the saved file.
    sheet.freeze_panes = "A2"

    header_cells = []
    for column in columns:
        cell = WriteOnlyCell(sheet, value=column.label_pl)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        header_cells.append(cell)
    sheet.append(header_cells)

    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.xlsx_width

    row_count = 0
    for record in records:
        row = []
        for column in columns:
            source = record.cell(column.key)
            value_text = source.parsed_value_text if source else None
            cell = WriteOnlyCell(sheet, value=None)

            if value_text is None or value_text == "":
                # Empty stays empty; it must never become a zero.
                row.append(cell)
                continue

            if column.type == ColumnType.MONEY:
                amount = _as_decimal(value_text)
                if amount is None:
                    cell.value = sanitize_text(str(value_text))
                else:
                    cell.value = amount
                    cell.number_format = column.xlsx_format or "0.00"
                    cell.alignment = Alignment(horizontal="right")
            elif column.type == ColumnType.DATE:
                date = _as_date(value_text)
                if date is None:
                    cell.value = sanitize_text(str(value_text))
                else:
                    cell.value = date
                    cell.number_format = column.xlsx_format or "yyyy-mm-dd"
                    cell.alignment = Alignment(horizontal="center")
            elif column.type == ColumnType.INTEGER:
                try:
                    cell.value = int(value_text)
                    cell.alignment = Alignment(horizontal="right")
                except ValueError:
                    cell.value = sanitize_text(str(value_text))
            else:
                cell.value = sanitize_text(str(value_text))
                cell.alignment = Alignment(wrap_text=True, vertical="top")

            if source is not None and source.is_manual:
                cell.font = Font(italic=True)
            row.append(cell)
        sheet.append(row)
        row_count += 1

    sheet.auto_filter.ref = f"A1:{get_column_letter(max(1, len(columns)))}{max(1, row_count + 1)}"

    if include_issues_sheet:
        _write_issues_sheet(workbook, profile, issues or [])
    if metadata:
        _write_metadata_sheet(workbook, metadata)

    workbook.save(target)
    workbook.close()

    data = target.read_bytes()
    return XlsxExportResult(
        path=target,
        row_count=row_count,
        column_count=len(columns),
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
    )


def _write_issues_sheet(workbook: Workbook, profile: KpirProfile, issues: Sequence[dict]) -> None:
    sheet = workbook.create_sheet(title=profile.issues_sheet_name[:31])
    headers = ["Kod", "Waga", "Strona", "Kolumna", "Szczegóły"]
    header_cells = []
    for text in headers:
        cell = WriteOnlyCell(sheet, value=text)
        cell.font = Font(bold=True)
        header_cells.append(cell)
    sheet.append(header_cells)
    for issue in issues:
        sheet.append(
            [
                sanitize_text(str(issue.get("code", ""))),
                sanitize_text(str(issue.get("severity", ""))),
                issue.get("page_number"),
                sanitize_text(str(issue.get("column_key", "") or "")),
                sanitize_text(str(issue.get("details", "") or "")),
            ]
        )
    for index, width in enumerate((26, 12, 8, 24, 60), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_metadata_sheet(workbook: Workbook, metadata: dict[str, str]) -> None:
    sheet = workbook.create_sheet(title="Metadane")
    header = []
    for text in ("Pole", "Wartość"):
        cell = WriteOnlyCell(sheet, value=text)
        cell.font = Font(bold=True)
        header.append(cell)
    sheet.append(header)
    for key, value in metadata.items():
        sheet.append([sanitize_text(str(key)), sanitize_text(str(value))])
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 60
